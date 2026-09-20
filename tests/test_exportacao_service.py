import pytest
from docx import Document
from io import BytesIO
from pypdf import PdfReader

from src.domain.peticao_schemas import Peca
from src.repositories.peca_repository import PecaRepository
from src.services import exportacao_service as exportacao_module
from src.services.exportacao_service import ExportacaoService


def _criar_peca(peca_texto: str = "Texto de teste.\n\nSegundo parágrafo.") -> Peca:
    return Peca(
        peca_id="peca-teste-1",
        template_id="consumo_indenizatoria_jec",
        titulo="Ação de Consumo",
        peca_texto=peca_texto,
    )


def test_gerar_docx_produz_arquivo_docx_valido():
    service = ExportacaoService()
    conteudo = service.gerar_docx(_criar_peca())

    assert conteudo[:2] == b"PK"  # docx é um zip
    doc = Document(BytesIO(conteudo))
    textos = [p.text for p in doc.paragraphs]
    assert any("Ação de Consumo" in t for t in textos)
    assert any("Texto de teste." in t for t in textos)


def test_gerar_pdf_produz_pdf_pesquisavel():
    service = ExportacaoService()
    conteudo = service.gerar_pdf(_criar_peca())

    assert conteudo[:4] == b"%PDF"
    reader = PdfReader(BytesIO(conteudo))
    texto_extraido = reader.pages[0].extract_text()
    assert "Ação de Consumo" in texto_extraido
    assert "Texto de teste." in texto_extraido


def test_pdf_com_caracteres_especiais_nao_quebra():
    """Travessão, aspas curvas e reticências unicode não podem derrubar a
    geração do PDF (fonte core só suporta latin-1)."""
    service = ExportacaoService()
    texto_dificil = "Fato — comprovado. Ele disse: \u201cnão pagarei\u201d… ponto final."
    conteudo = service.gerar_pdf(_criar_peca(texto_dificil))

    assert conteudo[:4] == b"%PDF"
    reader = PdfReader(BytesIO(conteudo))
    texto_extraido = reader.pages[0].extract_text()
    assert "comprovado" in texto_extraido
    assert "não pagarei" in texto_extraido  # acentuação preservada


def test_docx_preserva_travessao_e_aspas_curvas_sem_alteracao():
    """No docx (ao contrário do PDF) não há limitação de fonte, então o
    texto deve sair EXATAMENTE como veio, sem substituições."""
    service = ExportacaoService()
    texto_original = "Fato — comprovado com \u201caspas curvas\u201d."
    conteudo = service.gerar_docx(_criar_peca(texto_original))

    doc = Document(BytesIO(conteudo))
    textos = [p.text for p in doc.paragraphs]
    assert any(texto_original in t for t in textos)


def test_exportar_com_formato_invalido_levanta_erro(tmp_path, monkeypatch):
    repo = PecaRepository(db_path=str(tmp_path / "test.db"))
    repo.salvar(_criar_peca())
    monkeypatch.setattr(exportacao_module, "peca_repository", repo)

    service = ExportacaoService()
    with pytest.raises(ValueError):
        service.exportar("peca-teste-1", ["epub"])


def test_exportar_peca_inexistente_levanta_erro(tmp_path, monkeypatch):
    repo = PecaRepository(db_path=str(tmp_path / "test.db"))
    monkeypatch.setattr(exportacao_module, "peca_repository", repo)

    service = ExportacaoService()
    with pytest.raises(ValueError):
        service.exportar("nao-existe", ["docx"])


def test_exportar_ambos_formatos_retorna_os_dois(tmp_path, monkeypatch):
    repo = PecaRepository(db_path=str(tmp_path / "test.db"))
    repo.salvar(_criar_peca())
    monkeypatch.setattr(exportacao_module, "peca_repository", repo)

    service = ExportacaoService()
    resultado = service.exportar("peca-teste-1", ["docx", "pdf"])

    assert set(resultado.keys()) == {"docx", "pdf"}
    assert resultado["docx"][:2] == b"PK"
    assert resultado["pdf"][:4] == b"%PDF"
